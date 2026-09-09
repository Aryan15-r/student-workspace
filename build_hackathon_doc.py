from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'StudySpace_Hackathon_Documentation.docx'
NAVY = '173F5F'
BLUE = 'EAF3F8'
GRAY = 'D9E2F3'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}'); tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'4'); tag.set(qn('w:color'),'D9D9D9'); borders.append(tag)

def set_cell_text(cell, text, bold=False, color=None):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.font.color.rgb = RGBColor(0,0,0)
    return p

def para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); p.add_run(item)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); set_cell_text(c,h,True,'FFFFFF')
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if ri%2==1: shade(cells[i],BLUE)
            set_cell_text(cells[i],v)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal'].font.size=Pt(10.5)
for st in ['Heading 1','Heading 2']:
    styles[st].font.name='Aptos Display'; styles[st]._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); styles[st].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(16); styles['Heading 2'].font.size=Pt(12)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('StudySpace Hackathon Documentation'); r.font.name='Aptos Display'; r.font.size=Pt(25); r.font.bold=True; r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('A unified productivity workspace for students'); r.italic=True; r.font.size=Pt(13)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Hackathon Project Brief | Flutter, Supabase, and Gemini AI').font.size=Pt(10)

heading(doc,'Project Overview')
para(doc,'StudySpace is a cross-platform student productivity application that brings academic work into one focused workspace. Instead of moving between separate apps for assignments, doubts, discussion groups, resource discovery, scientific calculations, and document handling, a student can complete those workflows inside a single application.')
para(doc,'The project is designed for Android and the web from one Flutter codebase. It pairs a responsive client experience with Supabase for authentication, data, real-time community activity, and secure access rules, while Gemini powers the in-app AI assistant.')
table(doc,['Problem','StudySpace Response'],[
 ['Fragmented student workflows across many apps','A unified dashboard combines planning, AI help, communities, search, calculator, profile, and document tools.'],
 ['Assignments and exams are easy to overlook','The task planner stores priorities, categories, deadlines, completion state, and progress.'],
 ['Students need timely academic support','An AI assistant provides explanations and study support from inside the workspace.'],
 ['General chat groups become noisy','Community and channel structures keep academic discussion discoverable and organized.'],
 ['Students need tools even with weak connectivity','The scientific calculator is designed as an offline-capable utility.'],
], [2.25,4.7])

heading(doc,'Core Features')
table(doc,['Module','What it does','Student value'],[
 ['AI Assistant','Conversational study support with a Gemini API integration and formatted responses.','Helps students ask questions and get explanations without leaving the app.'],
 ['Task Planner','Creates and tracks assignments, exams, projects, personal tasks, deadlines, and priorities.','Turns scattered academic obligations into a clear work queue.'],
 ['Academic Community','Provides communities, channels, messages, reactions, and moderation reporting with real-time updates.','Creates a dedicated alternative to noisy class chat groups.'],
 ['Smart Search','Records and presents resource searches for learning material discovery.','Makes useful learning resources easier to find and revisit.'],
 ['Document and PDF Tools','Supports document selection, viewing, and PDF-oriented workflows in the application.','Reduces context switching to separate viewers and basic document utilities.'],
 ['Scientific Calculator','Evaluates math expressions locally and keeps a calculation workflow in the app.','Provides a quick study utility without an external app.'],
], [1.25,3.15,2.55])

heading(doc,'Technology Stack')
table(doc,['Layer','Technology','How it is used'],[
 ['Client application','Flutter and Dart','Single codebase for Android, iOS-ready architecture, and web/PWA delivery.'],
 ['State management','Provider','Separates screen state and business logic for tasks, auth, AI, community, dashboard, search, profile, calculator, and document tools.'],
 ['Navigation','go_router','Defines URL-aware, structured application routes.'],
 ['Backend platform','Supabase','Provides PostgreSQL, authentication, real-time messaging, storage, and security controls.'],
 ['Database','PostgreSQL','Stores profiles, tasks, communities, channels, messages, reactions, searches, and moderation reports.'],
 ['AI','Google Gemini API','Generates assistant responses through an HTTP-based service integration.'],
 ['UI and utility packages','Google Fonts, flutter_animate, shimmer, intl, timeago, uuid','Improve visual hierarchy, loading states, formatting, relative time, and identifier generation.'],
 ['Document utilities','pdf, pdfx, file_picker, open_filex, desktop_drop, archive','Support file selection, document and PDF creation/viewing, desktop interaction, and archive handling.'],
], [1.25,1.8,3.9])

heading(doc,'System Architecture')
para(doc,'The Flutter application is organized by feature. Presentation pages and widgets render the user experience; providers coordinate state; services and repositories connect features to Supabase, Gemini, and local utilities. This structure keeps each feature independently understandable while shared themes, widgets, exceptions, and helpers remain in core and shared folders.')
table(doc,['Flow','Implementation'],[
 ['User interaction','Flutter pages and reusable widgets collect input and present loading, empty, error, and success states.'],
 ['Feature state','Provider-backed state objects manage task lists, auth sessions, AI conversations, community data, and screen-specific actions.'],
 ['Data and identity','Supabase Auth identifies the user; PostgreSQL stores application records; Row Level Security limits data access.'],
 ['Live communication','Supabase Realtime propagates community message and reaction updates.'],
 ['AI generation','The AI service sends structured HTTP requests to Gemini and returns assistant output to the UI.'],
], [1.75,5.2])

heading(doc,'How We Built It')
para(doc,'We built StudySpace feature by feature around the student journey: sign in, see a dashboard, plan work, find support, collaborate, and use practical study tools. This kept the project scope focused on complete workflows rather than isolated screens.')
table(doc,['Step','What we did','Result'],[
 ['1. Defined the workspace','Mapped the recurring student needs: planning, doubt solving, resources, collaboration, and utility tools.','A focused set of modules that work together in one product.'],
 ['2. Created the Flutter foundation','Set up app routing, responsive scaffolding, theme tokens, shared widgets, and feature folders.','A consistent user experience and maintainable code organization.'],
 ['3. Added secure identity and data','Connected Supabase Auth and PostgreSQL migrations for profiles, tasks, communities, messages, reactions, searches, and reports.','Persistent user data with ownership-aware access patterns.'],
 ['4. Built student workflows','Implemented task cards and task creation, dashboard views, profile data, search, calculator, community channels, and document tools.','Practical tools usable from the same workspace.'],
 ['5. Integrated AI support','Added a Gemini-backed service and AI state layer so the assistant can respond inside the product.','Contextual academic assistance without leaving StudySpace.'],
 ['6. Hardened the experience','Used dedicated loading, error, and empty states; organized application exceptions; kept keys in environment configuration.','A clearer, safer experience when data or network operations fail.'],
], [0.6,3.5,2.85])

heading(doc,'Data Model and Security Basics')
para(doc,'The application extends Supabase Auth with a profiles table and connects each student-owned record to that identity. Tasks, search history, and profile edits are tied to the authenticated user. Community content is structured as communities, channels, messages, and reactions. Moderation reports provide a path for flagging inappropriate content.')
table(doc,['Security practice','Purpose'],[
 ['Row Level Security policies','Restrict user-owned records such as tasks and profiles to the authenticated owner.'],
 ['Environment configuration','Keeps Supabase and Gemini credentials out of source control and separates setup from code.'],
 ['Feature-level errors and loading states','Makes network and data failures visible and recoverable rather than silently failing.'],
 ['Structured database migrations','Makes the schema, indexes, seed communities, and policy setup repeatable across environments.'],
], [2.4,4.9])

heading(doc,'What Makes the Project Hackathon Ready')
bullets(doc,[
 'A demonstrable end-to-end product rather than a concept: users can authenticate, plan tasks, use study utilities, ask the assistant, and participate in channels.',
 'A clear real-world problem: students lose time and focus by switching among disconnected tools.',
 'A scalable foundation: Flutter supports multiple platforms, and Supabase provides hosted authentication, data, real-time updates, and policy controls.',
 'A modular codebase: each major student workflow has its own feature area, making future development and judging discussion easier.',
 'A practical demo path: show onboarding, create an urgent assignment, ask the AI for help, open a community channel, and demonstrate a calculator or document workflow.'
])

heading(doc,'Demo Script')
table(doc,['Time','Demo moment','Message to judges'],[
 ['0:00 - 0:20','Open the dashboard','StudySpace replaces app-switching with a single student workspace.'],
 ['0:20 - 0:50','Create an assignment with priority and deadline','Students can immediately organize academic work and see what matters next.'],
 ['0:50 - 1:20','Ask the AI assistant a course question','Help is available inside the same workflow rather than in a separate chatbot.'],
 ['1:20 - 1:50','Open a community channel and show live discussion','Course collaboration is organized by communities and channels, not an unstructured group chat.'],
 ['1:50 - 2:15','Use calculator or document tools','Common study utilities are accessible without opening another app.'],
 ['2:15 - 2:40','Explain the stack','Flutter enables cross-platform delivery; Supabase provides auth, data, real-time features, and security; Gemini powers AI support.'],
], [0.9,2.35,4.05])

heading(doc,'Future Scope')
bullets(doc,[
 'AI-generated flashcards, quizzes, and personalized study plans.',
 'Collaborative notes, whiteboards, and deeper group-study workflows.',
 'Expanded document processing such as smarter summaries, OCR, and conversion workflows.',
 'Campus-specific spaces, verified student communities, and institution-level deployment.',
 'Analytics that help students understand workload, upcoming deadlines, and study habits while respecting privacy.'
])

heading(doc,'Sources and Project References')
para(doc,'The following first-party project sources were used to document the implemented stack and architecture. External platform references are included for the technologies on which the project is built.')
table(doc,['Source','Use'],[
 ['Project README - github.com/Aryan15-r/student-workspace','Project overview, features, setup path, and high-level stack.'],
 ['pubspec.yaml in the project repository','Flutter package dependencies and client-side tooling.'],
 ['docs/schema.md and supabase/migrations','Database entities, indexes, realtime usage, and access-control design.'],
 ['lib/features and lib/app in the project repository','Feature-based architecture, presentation layers, providers, services, routing, and theme structure.'],
 ['https://flutter.dev','Flutter framework reference.'],
 ['https://supabase.com/docs','Supabase Auth, PostgreSQL, Realtime, and Row Level Security reference.'],
 ['https://ai.google.dev/gemini-api/docs','Gemini API integration reference.'],
 ['https://pub.dev','Dart and Flutter package documentation.'],
], [2.65,4.65])

heading(doc,'Quick Start Basics')
para(doc,'To run the project locally, clone the repository, create a .env file with the required Supabase and Gemini values, apply the supplied Supabase migration, run flutter pub get, and launch with flutter run. The repository README contains the full setup sequence and points to the schema documentation for the backend.')
para(doc,'Repository: https://github.com/Aryan15-r/student-workspace')
doc.save(OUT)
print(OUT)
