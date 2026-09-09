from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='StudySpace_Tech_Stack_PPT_Designer_Brief.docx'
NAVY='173F5F'; PALE='EAF3F8'

def shade(cell, fill):
    p=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:fill'),fill); p.append(s)
def cell(cell,text,bold=False,color=None):
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); r=p.add_run(text); r.bold=bold; r.font.size=Pt(9.2)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    b=OxmlElement('w:tcBorders')
    for e in ['top','left','bottom','right']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
    cell._tc.get_or_add_tcPr().append(b)
def title(doc,text,level=1):
    p=doc.add_paragraph(style='Heading '+str(level)); p.paragraph_format.space_before=Pt(13 if level==1 else 8); p.paragraph_format.space_after=Pt(5); p.add_run(text)
def p(doc,text,lead=None):
    z=doc.add_paragraph(); z.paragraph_format.space_after=Pt(5); z.paragraph_format.line_spacing=1.1
    if lead:
        r=z.add_run(lead);r.bold=True; z.add_run(text[len(lead):])
    else:z.add_run(text)
def bullets(doc,items):
    for x in items:
        z=doc.add_paragraph(style='List Bullet'); z.paragraph_format.space_after=Pt(2);z.add_run(x)
def tbl(doc,heads,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(heads):shade(t.rows[0].cells[i],NAVY);cell(t.rows[0].cells[i],x,True,'FFFFFF')
    for ri,row in enumerate(rows):
        for i,x in enumerate(row):
            c=t.add_row().cells[i]
            if ri%2:shade(c,PALE)
            cell(c,x)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

d=Document();s=d.sections[0];s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.7);s.right_margin=Inches(.7)
d.styles['Normal'].font.name='Aptos';d.styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos');d.styles['Normal'].font.size=Pt(10.2)
for k in ['Heading 1','Heading 2']:
 d.styles[k].font.name='Aptos Display';d.styles[k]._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display');d.styles[k].font.color.rgb=RGBColor(0,0,0)
d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 2'].font.size=Pt(11.5)
q=d.add_paragraph(style='Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=q.add_run('StudySpace Technical Stack');r.font.name='Aptos Display';r.font.size=Pt(25);r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run('PPT Designer Brief | Slide-ready architecture, technology, and proof points').italic=True

title(d,'How to Use This Brief')
p(d,'This document gives a presentation designer the exact technical story to visualize. Each numbered section can become one slide or a compact slide pair. The key visual language is a clean blue and white system: Flutter at the center, Supabase as the data and real-time backbone, Gemini as the AI layer, and student-facing modules surrounding them.')
tbl(d,['Slide','Recommended visual','Key message'],[
['1','Three-layer architecture diagram','StudySpace combines a cross-platform Flutter client, a managed cloud backend, and an AI intelligence layer.'],
['2','Technology logo wall or stack cards','Every technology has a defined responsibility; this is not a collection of random tools.'],
['3','Feature-to-technology matrix','Eight student workflows are implemented through reusable application modules.'],
['4','Data flow diagram','A user action moves through UI, Provider state, service/repository logic, then Supabase or Gemini.'],
['5','Database ER diagram','A structured relational model supports identity, planning, collaboration, search, and moderation.'],
['6','Security shield diagram','Authentication, Row Level Security, environment configuration, and scoped access protect user data.'],
['7','AI pipeline','Gemini model fallback and an offline smart-response path keep the assistant responsive.'],
['8','Developer architecture diagram','Feature-first folders, routing, shared design tokens, and focused providers make the app maintainable.'],
],[.55,2.9,3.8])

title(d,'1. Top-Level Architecture')
p(d,'Designer direction: use a three-tier stack with arrows flowing vertically. Place the StudySpace logo or title above the diagram. Avoid describing this as a conventional server-heavy backend: Supabase supplies managed backend capabilities, while Flutter owns the app experience.')
tbl(d,['Layer','Technologies','Responsibility','Visual treatment'],[
['Presentation layer','Flutter, Dart, Material UI, Google Fonts','Renders responsive Android and web experiences; contains pages, widgets, themed navigation, loading/error/empty states.','Top row: mobile and browser mockups.'],
['Application layer','Provider, go_router, feature services, repositories','Coordinates state, navigation, business logic, error mapping, and API/database calls.','Middle row: connected logic blocks or a routing map.'],
['Cloud and intelligence layer','Supabase, PostgreSQL, Auth, Realtime, Gemini API','Stores data, authenticates students, powers live discussion, enforces policies, and generates AI responses.','Bottom row: database, lock, lightning, and AI icons.'],
],[1.15,2.15,2.65,1.25])
p(d,'Suggested diagram labels: Student -> Flutter UI -> Provider / Service / Repository -> Supabase Auth and Database or Gemini AI -> updated UI. For community chat, add a return arrow from Supabase Realtime to the Flutter UI to show live updates.')

title(d,'2. Core Technology Stack')
tbl(d,['Category','Technology','Why it is in the stack','PPT visual cue'],[
['Cross-platform app','Flutter and Dart','One codebase targets Android and web/PWA and can extend to iOS. Flutter provides the UI framework; Dart powers application logic.','Flutter logo + Android phone + browser.'],
['State management','Provider','Keeps UI state separate from data and business logic. Providers exist for auth, dashboard, tasks, AI, communities, search, profile, calculator, PDF tools, and study tools.','State nodes feeding screens.'],
['Navigation','go_router','Centralized, URL-aware routing with auth-aware redirects and nested community/channel routes.','Route map / signpost.'],
['Backend platform','Supabase','Managed backend platform used for authentication, PostgreSQL data, realtime messaging, storage-oriented capabilities, and access policies.','Supabase logo + cloud.'],
['Database','PostgreSQL','Relational persistence for profiles, tasks, communities, channels, messages, reactions, searches, and reports.','Database cylinder with linked tables.'],
['Authentication','Supabase Auth','Supports the signed-in identity used to associate users with profiles, tasks, messages, reactions, and reports.','Shield + user icon.'],
['Live updates','Supabase Realtime','Publishes message and reaction changes to keep community conversations current.','Lightning / broadcast waves.'],
['Artificial intelligence','Google Gemini REST API','Generates AI assistant replies and structured academic-search data.','Gemini star + chat bubbles.'],
['Networking and config','http, flutter_dotenv','Sends API requests and loads keys such as GEMINI_API_KEY from environment configuration.','API arrow + key icon.'],
],[1.05,1.45,3.65,1.05])

title(d,'3. Feature-to-Stack Map')
tbl(d,['Student capability','Frontend implementation','Backend or intelligence dependency','Outcome'],[
['Account access','Auth pages, AuthProvider, protected routes','Supabase Auth and profiles','A student has a persistent identity and personalized workspace.'],
['Dashboard','Dashboard page and provider','Supabase-backed feature data','A single landing point for student activity.'],
['Task planner','Todo page, task cards, task provider, TodoRepository','tasks table in PostgreSQL','Assignments, exams, projects, and personal work are tracked with priorities and dates.'],
['AI assistant','AI page, chat models, AI provider, AiService','Gemini generateContent API; local fallback','Academic questions are handled inside the product.'],
['Academic search','Search page and provider','Gemini JSON search mode; search_history','Search overview, resource results, and related topics can be shown in a structured UI.'],
['Communities','Community, channel, and chat pages; CommunityProvider','communities, channels, messages, reactions; Realtime','Students exchange course-focused messages in organized channels.'],
['Document and PDF tools','PdfToolsPage, file picker, embedded viewer, PDF packages','Local file handling and document packages','Students can select, view, extract, and work with study documents.'],
['Calculator and focus tools','Calculator page and provider; math_expressions','Local Dart computation','A useful offline-friendly study utility is available without a separate app.'],
],[1.2,2.0,2.1,1.9])

title(d,'4. Application Architecture and Code Organization')
p(d,'StudySpace follows a feature-first Flutter structure. The designer should visualize the project as repeatable vertical slices rather than one large block of code. Each major student capability is grouped under lib/features, while app-wide routing and theme live under lib/app and shared UI/utilities live under lib/core and lib/shared.')
tbl(d,['Folder area','Role','Examples to name on slide'],[
['lib/app','Application shell','app.dart, router.dart, app_theme.dart, app_colors.dart, app_text_styles.dart'],
['lib/core','Cross-cutting functionality','errors, file-saving utilities, loading, empty, and error widgets'],
['lib/features','Feature modules','auth, dashboard, todo, ai_assistant, search, community, calculator, pdf_tools, profile, study_tools'],
['lib/shared','Reusable UI','adaptive scaffold, app logo, login prompt dialog'],
['supabase/migrations','Repeatable backend schema','initial schema, channel policy improvements, private rooms and moderation, task times'],
],[1.45,2.2,3.55])
bullets(d,[
'Recommended code-flow visual: Page/Widget -> Provider -> Service or Repository -> Supabase or Gemini -> Provider state -> updated Page/Widget.',
'Routing is centralized in AppRouter. It redirects unauthenticated users away from protected features while leaving selected focus tools available without login.',
'Dedicated shared components for loading, errors, and empty states make asynchronous features consistent and easier to present as production-minded engineering.',
])

title(d,'5. Data Architecture')
p(d,'Designer direction: use an entity-relationship diagram. Put auth.users at the upper left, connected one-to-one to profiles. Place tasks and search_history below profiles as user-owned data. Place communities -> channels -> messages -> message_reactions as a separate collaboration chain. Attach reports to messages and the reporting user.')
tbl(d,['Entity','Important fields or relation','Why it exists'],[
['profiles','Primary key matches auth.users.id; username, name, avatar, college, branch, year','Extends authentication identity into a student profile.'],
['tasks','user_id, title, category, priority, due_date, completed','Stores a student-owned academic and personal task list.'],
['communities','name, description, icon, category, created_by','Creates topic-level discussion spaces such as Programming or Mathematics.'],
['channels','community_id, name, description','Organizes discussion inside a community, such as #general or homework help.'],
['messages','channel_id, user_id, content, timestamps','Stores channel conversations; configured for Realtime publication.'],
['message_reactions','message_id, user_id, emoji; unique composite rule','Lets a user react once per emoji per message without duplicate reactions.'],
['search_history','user_id, query, created_at','Supports recent searches and personalization.'],
['reports','reporter_id, message_id, reason','Provides a moderation reporting workflow for community content.'],
],[1.35,3.4,2.45])
p(d,'Performance details for a small secondary callout: indexes exist for common lookups including task user and due date, messages by channel and descending time, channels by community, search history by user, and reactions by message.')

title(d,'6. Authentication and Security Story')
tbl(d,['Control','Implemented approach','Designer-friendly explanation'],[
['Identity','Supabase Auth with a profile-creation database trigger','When a student signs up, the backend automatically creates the matching profile record.'],
['Access control','PostgreSQL Row Level Security policies','A student can access their own tasks, search history, and account updates; message mutation is tied to its author.'],
['Data ownership','user_id foreign keys linked to profiles','Every personal action has a clear owner in the database.'],
['API credential handling','flutter_dotenv loads configuration from .env','Secrets are configured outside the application source rather than hardcoded in screens.'],
['Error handling','AppException mapping and dedicated UI states','Network, database, or API failures have a predictable user-facing path.'],
['Input and data consistency','PostgreSQL types, checks, foreign keys, unique constraints, and migrations','The database guards category values, priority values, relations, and duplicate reactions.'],
],[1.35,3.25,2.6])
p(d,'Important accuracy note for the designer: communities, channels, profiles, messages, and reactions are readable under the current schema policies; user-owned writes and changes are guarded by authentication and ownership checks. Do not claim end-to-end encrypted chat or claim that all data is private.')

title(d,'7. AI Architecture')
p(d,'The AI assistant uses Google Gemini through a REST integration. The application builds a conversation payload from message history, adds a system instruction, and sends it with a temperature of 0.7 and a maximum output allowance of 16,384 tokens. The app tries a small ordered model list: gemini-2.5-flash, gemini-2.0-flash, then gemini-flash-latest. If a request fails or no key is configured, it returns a local academic fallback response rather than leaving the user at a dead end.')
tbl(d,['AI stage','Technical behavior','Visual cue'],[
['1. Capture context','Chat history and the newest user message are encoded as Gemini content parts.','Conversation bubbles entering pipeline.'],
['2. Secure configuration','The Gemini key is read from GEMINI_API_KEY in environment configuration.','Key icon before API gateway.'],
['3. Request generation','HTTP POST targets Gemini generateContent with system instruction, safety settings, temperature, and token limit.','API request arrow.'],
['4. Resilience','The app tries multiple Flash model identifiers in order and applies timeouts.','Three linked AI nodes / fallback arrow.'],
['5. Response quality','All response parts are joined; mathematical formatting is cleaned before display.','Raw response -> polished study answer.'],
['6. Fallback','When API access is unavailable, the app produces a local structured academic response.','Offline mode / safety net.'],
['7. Search mode','A lower-temperature request asks Gemini for raw JSON: overview, resources, related topics, source metadata, and free/paid status.','Structured JSON brackets.'],
],[1.35,4.25,1.65])

title(d,'8. Document Tools and Local Utility Stack')
tbl(d,['Package or approach','Role in the product','Slide wording'],[
['file_picker and desktop_drop','Accepts files from device selection and desktop drag/drop workflows.','Bring study documents into the workspace.'],
['pdf and pdfx','Supports PDF generation, viewing, and pinch/zoom viewing flows.','Read and work with PDFs in-app.'],
['flutter_markdown and markdown','Renders structured content such as AI answers in a readable format.','Present rich AI answers cleanly.'],
['math_expressions','Parses and evaluates calculator expressions locally.','A scientific calculation tool that does not depend on a remote database.'],
['shared_preferences','Stores suitable lightweight local preferences.','Keeps small client-side preferences available.'],
['path_provider, open_filex, archive','Supports file locations, opening local outputs, and archive-oriented file processing.','Utility foundation for local document workflows.'],
],[1.65,3.65,1.95])

title(d,'9. PPT Design Checklist and Approved Claims')
bullets(d,[
'Use platform logos only for Flutter, Supabase, PostgreSQL, Gemini, Dart, Provider, and go_router. Use generic icons for the package-level utilities to avoid a noisy logo wall.',
'Color direction: deep navy #173F5F for headings and data layers, white background, pale blue #EAF3F8 table/diagram surfaces, with one bright accent only if needed.',
'Best hero diagram: Flutter UI at the top; Provider and feature services in the middle; Supabase Auth, PostgreSQL, Realtime, and Gemini as bottom services.',
'Best proof-point line: One Flutter codebase connects planning, AI support, academic discovery, real-time communities, and document utilities for students.',
'Use these exact defensible claims: cross-platform Flutter app; Supabase-backed authentication and PostgreSQL; Realtime-enabled messages and reactions; Row Level Security; Gemini-powered assistant with fallback behavior; feature-first modular architecture.',
'Avoid unsupported claims: guaranteed zero downtime, 100 percent uptime, end-to-end encryption, OCR if not demonstrated, native editing of every Office format, or a production-scale user count.',
])
title(d,'Source Files for Designer and Presenter')
tbl(d,['Repository source','What it verifies'],[
['README.md','Product positioning, core features, baseline technical stack, and local setup.'],
['pubspec.yaml','Flutter package dependencies used in the build.'],
['lib/app/router.dart','Routes, authentication redirect rules, and product navigation surface.'],
['lib/features/ai_assistant/services/ai_service.dart','Gemini API payloads, models, timeouts, response processing, search JSON mode, and fallback behavior.'],
['lib/features/todo/data/todo_repository.dart','Supabase task CRUD and app exception handling pattern.'],
['docs/schema.md and supabase/migrations/001_initial_schema.sql','Entities, policies, indexes, triggers, default communities, and Realtime publication.'],
['lib/features/pdf_tools/presentation/pages/pdf_tools_page.dart','In-app document and PDF interaction implementation.'],
],[2.65,4.6])
d.save(OUT);print(OUT)
